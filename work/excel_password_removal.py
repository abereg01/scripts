from openpyxl import load_workbook

def unprotect_excel(filename):
    wb = load_workbook(filename)
    for sheet in wb.sheetnames:
        wb[sheet].protection.sheet = False
    wb.save('unprotected_' + filename)

# Usage
unprotect_excel('1-mall-ersattning-samforlaggning_2021.xlsx')
